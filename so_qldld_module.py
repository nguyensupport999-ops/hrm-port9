# -*- coding: utf-8 -*-
"""
Module: Sổ quản lý lao động (theo Nghị định 145/2020/NĐ-CP, mẫu ban hành kèm
Thông tư hướng dẫn của Bộ LĐTBXH).

Gồm 2 phần:
  1. build_so_qldld_excel(...)   -> hàm điền dữ liệu vào file mẫu Excel
  2. render_tab_so_qldld(...)    -> code Streamlit cho tab t4, dán trực tiếp
     vào khối "with t4:" trong menu 📋 BHXH.

CÁCH TÍCH HỢP VÀO FILE APP HIỆN TẠI
------------------------------------
1) Đổi dòng khai báo tab từ:
       t1, t2, t3 = st.tabs(["📊 Tổng quan", "📝 Báo cáo tăng/giảm D02-LT", "📥 Xuất BC trích nộp BH"])
   thành:
       t1, t2, t3, t4 = st.tabs([
           "📊 Tổng quan", "📝 Báo cáo tăng/giảm D02-LT",
           "📥 Xuất BC trích nộp BH", "📔 Sổ quản lý lao động",
       ])

2) Thêm ở cuối file (sau khối "with t3: render_xuat_bao_cao_bhxh(...)"):
       with t4:
           render_tab_so_qldld(
               st.session_state.db_engine,
               format_date=format_date,
               company_config=COMPANY_CONFIG,
               auto_download_excel=_auto_download_excel,
           )
   (3 tham số format_date/company_config/auto_download_excel đã tồn tại sẵn
   trong app.py -> TRUYỀN VÀO, KHÔNG import ngược lại từ app.py, vì Streamlit
   chạy app.py dưới tên module "__main__" chứ không phải "app" -> nếu import
   ngược sẽ vô tình nạp lại và chạy lại toàn bộ app.py từ đầu, gây lỗi.)

3) Copy file mẫu gốc (đính kèm) vào:
       excel_templates/SoQLLD_template_goc.xlsx
   (giữ nguyên tên/khớp với TEMPLATE_SO_QLLD bên dưới, hoặc sửa lại đường dẫn).

4) import các hàm bên dưới vào file app chính:
       from so_qldld_module import build_so_qldld_excel, render_tab_so_qldld
"""

import os
from copy import copy as _copy
from datetime import date

import openpyxl
import openpyxl.styles
from openpyxl.utils import get_column_letter, column_index_from_string

TEMPLATE_SO_QLLD = "excel_templates/SoQLLD_template_goc.xlsx"

# ============================================================================
# VỊ TRÍ DÒNG/CỘT CỐ ĐỊNH - đã xác minh trực tiếp trên file mẫu thật
# (SoQLLD_template_goc.xlsx). Nếu sau này thay mẫu khác, kiểm tra lại các
# hằng số này cho khớp.
# ============================================================================
HEADER_ROW = 5          # Dòng tiêu đề cột (merge dọc 2 dòng cho hầu hết các cột)
SUB_HEADER_ROW = 6      # Dòng tiêu đề phụ (chỉ 3 cột con BHXH/BHYT/BHTN)
FIRST_DATA_ROW = 7      # Dòng dữ liệu lao động đầu tiên

# Cột "Mã NV" được CHÈN THÊM ngay sau cột STT -> chèn tại vị trí cột B (2),
# đẩy toàn bộ các cột còn lại (Họ tên, Giới tính, ... Chấm dứt HĐLĐ) sang
# phải 1 cột so với mẫu gốc.
MA_NV_INSERT_AT = 2

# ----- Chỉ số cột SAU KHI ĐÃ CHÈN xong cột Mã NV (dùng để ghi dữ liệu) -----
COL_STT = 1
COL_MA_NV = 2
COL_HO_TEN = 3
COL_GIOI_TINH = 4
COL_NGAY_SINH = 5
COL_QUOC_TICH = 6
COL_NOI_CU_TRU = 7            # = cột "F" trong mẫu gốc (trước khi chèn Mã NV)
COL_SO_CCCD = 8
COL_TRINH_DO = 9
COL_BAC_TRINH_DO_NGHE = 10
COL_VI_TRI_LAM_VIEC = 11
COL_LOAI_HDLD = 12
COL_BAT_DAU_LAM_VIEC = 13
COL_BHXH = 14
COL_BHYT = 15
COL_BHTN = 16
COL_TIEN_LUONG = 17           # = cột "P" trong mẫu gốc (trước khi chèn Mã NV)
# Cột 18-24 (mẫu gốc là Q..W): Nâng bậc lương, Số ngày nghỉ, Giờ làm thêm,
# Hưởng chế độ BHXH/BHYT/BHTN, Học nghề đào tạo, Kỷ luật lao động, TNLĐ-BNN
# -> hệ thống hiện chưa lưu các sự kiện này dưới dạng trường riêng, để trống,
# cập nhật thủ công theo phát sinh thực tế trong kỳ.
COL_CHAM_DUT_HDLD = 25        # = cột "X" trong mẫu gốc (trước khi chèn Mã NV)
N_COLS = 25                   # Tổng số cột sau khi chèn (A..Y)

# ----- Thứ tự nhóm bắt buộc theo yêu cầu -----
GROUP_LABELS = [
    "Không xác định thời hạn",
    "Hợp đồng có thời hạn",
    "Thử việc",
    "Khác",                              # nhóm dự phòng, chỉ xuất hiện nếu có
    "Đã chấm dứt hợp đồng lao động",      # luôn ở cuối cùng
]
GRP_KHONG_XD = 0
GRP_CO_THOI_HAN = 1
GRP_THU_VIEC = 2
GRP_KHAC = 3
GRP_CHAM_DUT = 4


def _phan_nhom(nv):
    """Xác định nhóm của 1 lao động theo đúng thứ tự yêu cầu:
    1) Không xác định thời hạn
    2) Hợp đồng có thời hạn  (DB lưu là "Xác định thời hạn", kể cả biến thể
       "Xác định thời hạn 12/24/36 tháng")
    3) Thử việc
    4) (dự phòng) Khác - loai_hop_dong không khớp mẫu nào ở trên
    5) Đã chấm dứt hợp đồng lao động - ĐƯỢC ƯU TIÊN CAO NHẤT, dựa vào
       trang_thai == 'NGHI_VIEC' (KHÔNG dựa vào ngay_ket_thuc khác NULL nữa,
       vì với lao động "Thử việc", ngay_ket_thuc có thể đang lưu NGÀY HẾT
       HẠN THỬ VIỆC DỰ KIẾN chứ không phải ngày họ thực sự nghỉ -> nếu dùng
       ngay_ket_thuc sẽ vô tình xếp nhầm cả người đang thử việc bình thường
       vào nhóm "đã chấm dứt". trang_thai là cờ trạng thái thực tế, đáng tin
       cậy hơn).
    """
    if (nv.get("trang_thai") or "").strip().upper() == "NGHI_VIEC":
        return GRP_CHAM_DUT
    lhd = (nv.get("loai_hop_dong") or "").strip().lower()
    if lhd == "không xác định thời hạn":
        return GRP_KHONG_XD
    if "xác định thời hạn" in lhd:  # khớp cả "xác định thời hạn 12 tháng"...
        return GRP_CO_THOI_HAN
    if lhd == "thử việc":
        return GRP_THU_VIEC
    return GRP_KHAC


def _to_number(val):
    """Chuyển giá trị lương (có thể là str '7000000', '4650000.0', số, hoặc
    None) về kiểu số THẬT (int/float). Nếu ghi giá trị dạng chuỗi vào ô Excel,
    number_format '#,##0' sẽ KHÔNG được áp dụng (Excel chỉ định dạng số cho
    kiểu numeric) -> đây là gốc rễ của lỗi "cột P chưa định dạng được"."""
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return val
    try:
        s = str(val).strip().replace(",", "")
        if s == "":
            return None
        f = float(s)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return None


def _insert_column_after(ws, insert_col, max_row=None, max_col=None):
    """Chèn 1 cột trống vào vị trí insert_col, dịch toàn bộ nội dung / style /
    độ rộng cột / vùng merge từ insert_col trở đi sang phải 1 cột.

    QUAN TRỌNG: KHÔNG dùng ws.insert_cols() của openpyxl ở đây - đã kiểm
    chứng trực tiếp trên file mẫu thật rằng nó dịch chuyển GIÁ TRỊ Ô đúng
    nhưng KHÔNG dịch các VÙNG MERGE (merged_cells) một cách tương ứng, làm
    lệch toàn bộ header 2 dòng của mẫu (ví dụ nhóm merge "Tham gia bảo hiểm"
    bị lệch khỏi 3 cột con BHXH/BHYT/BHTN, cột cuối cùng mất merge). Hàm này
    tự làm đúng cả 2 việc: dịch nội dung/style VÀ dịch/nới merge.
    """
    if max_row is None:
        max_row = ws.max_row
    if max_col is None:
        max_col = ws.max_column

    # 1) Gỡ toàn bộ merge hiện có trước khi đụng vào cell (openpyxl không
    #    cho set value trực tiếp vào ô không phải góc trên-trái của 1 vùng
    #    đang merge).
    old_merges = list(ws.merged_cells.ranges)
    for m in old_merges:
        ws.unmerge_cells(str(m))

    new_merges = []
    for m in old_merges:
        min_col, min_row, max_c, max_r = m.min_col, m.min_row, m.max_col, m.max_row
        if min_col >= insert_col:
            # Toàn bộ vùng merge nằm sau điểm chèn -> dịch cả 2 mép phải 1 cột
            min_col += 1
            max_c += 1
        elif max_c >= insert_col:
            # Vùng merge "bao trùm" điểm chèn (vd. A1:F1 tiêu đề DN) -> nới
            # rộng thêm 1 cột để vẫn phủ đúng độ rộng bảng như trước khi chèn
            max_c += 1
        new_merges.append((min_row, min_col, max_r, max_c))

    # 2) Dịch nội dung + style từng ô, DUYỆT TỪ PHẢI SANG TRÁI để không ghi
    #    đè lên dữ liệu chưa kịp dịch.
    for col in range(max_col, insert_col - 1, -1):
        for row in range(1, max_row + 1):
            src = ws.cell(row=row, column=col)
            dst = ws.cell(row=row, column=col + 1)
            dst.value = src.value
            dst.font = _copy(src.font)
            dst.border = _copy(src.border)
            dst.fill = _copy(src.fill)
            dst.alignment = _copy(src.alignment)
            dst.number_format = src.number_format

    # 3) Làm trống cột vừa "nhường chỗ" (insert_col), lấy tạm style của cột
    #    kế bên (đã dịch) để viền/font đồng bộ với các cột xung quanh.
    for row in range(1, max_row + 1):
        c = ws.cell(row=row, column=insert_col)
        c.value = None
        neighbor = ws.cell(row=row, column=insert_col + 1)
        c.font = _copy(neighbor.font)
        c.border = _copy(neighbor.border)
        c.fill = _copy(neighbor.fill)
        c.alignment = _copy(neighbor.alignment)
        c.number_format = neighbor.number_format

    # 4) Áp lại các vùng merge đã dịch/nới ở bước 1.
    for row1, col1, row2, col2 in new_merges:
        ws.merge_cells(start_row=row1, start_column=col1, end_row=row2, end_column=col2)

    # 5) Dịch độ rộng cột (column_dimensions), key theo chữ cái -> duyệt từ
    #    cột lớn nhất về nhỏ nhất để không ghi đè khi shift.
    old_widths = {}
    for letter, dim in list(ws.column_dimensions.items()):
        if dim.width is not None:
            old_widths[column_index_from_string(letter)] = dim.width
    for idx in sorted(old_widths.keys(), reverse=True):
        if idx >= insert_col:
            ws.column_dimensions[get_column_letter(idx + 1)].width = old_widths[idx]
    # Cột mới chèn (Mã NV): đặt độ rộng mặc định vừa phải.
    ws.column_dimensions[get_column_letter(insert_col)].width = 10


def _chen_cot_ma_nv(ws):
    """Chèn cột 'Mã NV' ngay sau cột STT, rồi ghi tiêu đề cho cột mới (merge
    dọc 2 dòng HEADER_ROW:SUB_HEADER_ROW giống các cột đơn khác trong mẫu)."""
    _insert_column_after(ws, MA_NV_INSERT_AT)

    header_style = ws.cell(row=HEADER_ROW, column=COL_STT)
    cell = ws.cell(row=HEADER_ROW, column=COL_MA_NV, value="Mã NV")
    cell.font = _copy(header_style.font)
    cell.border = _copy(header_style.border)
    cell.fill = _copy(header_style.fill)
    cell.alignment = _copy(header_style.alignment)
    ws.merge_cells(
        start_row=HEADER_ROW, start_column=COL_MA_NV,
        end_row=SUB_HEADER_ROW, end_column=COL_MA_NV,
    )


def _dong_bo_do_rong_cot(ws):
    """Set độ rộng cột 'Nơi cư trú' và cột 'Chấm dứt HĐLĐ và lý do' bằng
    nhau (lấy theo độ rộng đã có sẵn của cột 'Nơi cư trú' trong mẫu, vì cột
    này vốn đã được set đủ rộng để chứa địa chỉ dài)."""
    target_width = ws.column_dimensions[get_column_letter(COL_NOI_CU_TRU)].width
    if target_width:
        ws.column_dimensions[get_column_letter(COL_CHAM_DUT_HDLD)].width = target_width


def _write_group_header(ws, row, label, n_cols=N_COLS):
    """Ghi 1 dòng tiêu đề nhóm, merge từ cột A đến cột cuối bảng, in đậm,
    CĂN TRÁI CÓ THỤT LỀ (indent), căn giữa theo chiều dọc, nền xám nhạt để
    phân biệt với dòng dữ liệu."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    ws.row_dimensions[row].height = 20


def _chars_per_line(ws, col_idx, fallback=28):
    """Ước lượng số ký tự vừa 1 dòng trong cột, dựa vào độ rộng cột đã đặt
    trong file mẫu (nếu có), dùng để tính chiều cao dòng khi wrap text."""
    dim = ws.column_dimensions.get(get_column_letter(col_idx))
    if dim and dim.width:
        return max(8, int(dim.width) - 2)
    return fallback


def _autosize_row_for_wrap(ws, row, col, text, line_height=15, min_height=15):
    """Tăng chiều cao dòng (nếu cần) để text wrap trong 1 ô hiển thị đủ."""
    if not text:
        return
    chars_per_line = _chars_per_line(ws, col)
    n_lines = max(1, -(-len(str(text)) // chars_per_line))  # ceil division
    needed = max(min_height, n_lines * line_height)
    current = ws.row_dimensions[row].height or min_height
    if needed > current:
        ws.row_dimensions[row].height = needed


def _gioi_tinh_display(gt):
    if not gt:
        return ""
    gt = str(gt).strip().lower()
    if gt in ("nam", "m", "male"):
        return "Nam"
    if gt in ("nữ", "nu", "f", "female"):
        return "Nữ"
    return gt


def _fmt_ngay(d):
    if not d:
        return ""
    if isinstance(d, str):
        return d
    return d.strftime("%d/%m/%Y")


def _capture_template_row_style(ws, src_row, n_cols=N_COLS):
    """Chụp (snapshot) style của dòng mẫu (font/border/fill/alignment/
    number_format + chiều cao dòng) vào bộ nhớ.

    QUAN TRỌNG: phải gọi hàm này SAU KHI đã chèn xong cột Mã NV (để bắt đúng
    style của dòng mẫu với đủ N_COLS cột), và TRƯỚC KHI ghi bất cứ nội dung
    nào (kể cả tiêu đề nhóm) đè lên dòng mẫu (FIRST_DATA_ROW). Vì việc phân
    nhóm có thể khiến chính dòng FIRST_DATA_ROW trở thành dòng tiêu đề nhóm
    đầu tiên (nếu nhóm đó có lao động) -> nếu không chụp trước, các dòng dữ
    liệu phía sau sẽ vô tình copy nhầm định dạng "tiêu đề nhóm" (in đậm, nền
    xám, merge) thay vì định dạng dòng dữ liệu bình thường.
    """
    styles = []
    for col in range(1, n_cols + 1):
        src_cell = ws.cell(row=src_row, column=col)
        styles.append({
            "font": _copy(src_cell.font),
            "border": _copy(src_cell.border),
            "fill": _copy(src_cell.fill),
            "alignment": _copy(src_cell.alignment),
            "number_format": src_cell.number_format,
        })
    row_height = ws.row_dimensions[src_row].height
    return {"cols": styles, "row_height": row_height}


def _apply_template_row_style(ws, dst_row, template_style, n_cols=N_COLS):
    """Áp style đã chụp (từ _capture_template_row_style) vào dst_row."""
    cols_style = template_style["cols"]
    for col in range(1, n_cols + 1):
        if col > len(cols_style):
            continue
        style = cols_style[col - 1]
        dst_cell = ws.cell(row=dst_row, column=col)
        dst_cell.font = style["font"]
        dst_cell.border = style["border"]
        dst_cell.fill = style["fill"]
        dst_cell.alignment = style["alignment"]
        dst_cell.number_format = style["number_format"]
    if template_style.get("row_height"):
        ws.row_dimensions[dst_row].height = template_style["row_height"]


def build_so_qldld_excel(template_path, output_path, employees, ngay_bao_cao, company_config):
    """
    Điền danh sách lao động vào file mẫu "Sổ quản lý lao động".

    employees: list[dict] - mỗi dict là 1 bản ghi lấy từ bảng nhan_vien, cần các
        khoá (lấy None/"" nếu không có, hàm tự bỏ qua):
        ma_nv, ho_ten, gioi_tinh, ngay_sinh, quoc_tich, thuong_tru, so_cccd,
        trinh_do, bac_trinh_do_nghe, chuc_danh_nghe, loai_hop_dong,
        ngay_vao_lam, thang_bat_dau_bh, luong_bao_hiem, ngay_ket_thuc, ly_do_nghi
    ngay_bao_cao: date - mốc thời điểm "chụp" trạng thái sổ (in vào tiêu đề sổ).
        LƯU Ý: Sổ quản lý lao động theo Điều 3 Nghị định 145/2020/NĐ-CP là một
        SỔ TÍCH LŨY (ghi nhận toàn bộ lịch sử lao động của doanh nghiệp, cập
        nhật liên tục "kể từ ngày người lao động bắt đầu làm việc"), KHÔNG
        phải một báo cáo biến động theo kỳ như Điều 4 (Báo cáo sử dụng lao
        động, Mẫu 01/PLI). Vì vậy hàm này chỉ nhận 1 mốc "tính đến ngày", chứ
        không nhận khoảng "từ ngày - đến ngày".
    company_config: dict - {"ten_cong_ty":..., "ma_so_thue":..., "dia_chi":...}

    Danh sách được TỰ ĐỘNG PHÂN NHÓM và in theo đúng thứ tự:
      1) Không xác định thời hạn
      2) Hợp đồng có thời hạn
      3) Thử việc
      4) (nếu có phát sinh) Khác
      5) Đã chấm dứt hợp đồng lao động - luôn ở cuối (dựa vào trang_thai).
    Mỗi nhóm có 1 dòng tiêu đề (merge toàn bảng, in đậm, căn trái có thụt lề,
    kèm số lượng, ví dụ "Không xác định thời hạn: 65 người"). Nhóm không có
    lao động nào thì không in dòng tiêu đề của nhóm đó. Trong mỗi nhóm, lao
    động được sắp xếp theo ma_nv tăng dần. STT đánh số liên tục xuyên suốt
    toàn bộ sổ (không reset lại theo từng nhóm).

    Cột "Mã NV" được chèn ngay sau cột STT (đẩy các cột còn lại sang phải).
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Sheet1"]

    # ----- Chèn cột "Mã NV" ngay sau STT (phải làm TRƯỚC mọi bước khác vì
    # nó dịch chuyển toàn bộ vị trí cột phía sau) -----
    _chen_cot_ma_nv(ws)

    # ----- Đồng bộ độ rộng cột "Nơi cư trú" và "Chấm dứt HĐLĐ và lý do" -----
    _dong_bo_do_rong_cot(ws)

    # ----- Tiêu đề doanh nghiệp -----
    # Ưu tiên khoá đúng trong COMPANY_CONFIG thật của app ("ten_cong_ty",
    # "ma_so_thue"); vẫn giữ khoá cũ ("ten_doanh_nghiep", "mst") làm dự phòng
    # để không vỡ nếu app dùng bộ khoá khác.
    ten_cong_ty = company_config.get("ten_cong_ty") or company_config.get("ten_doanh_nghiep", "")
    ma_so_thue = company_config.get("ma_so_thue") or company_config.get("mst", "")
    ws["A1"] = f"DOANH NGHIỆP: {ten_cong_ty}"
    ws["A2"] = f"Mã số thuế: {ma_so_thue}"
    ws["A3"] = f"Địa chỉ: {company_config.get('dia_chi', '')}"
    ws["A4"] = f"SỔ QUẢN LÝ LAO ĐỘNG (Tính đến ngày {ngay_bao_cao.strftime('%d/%m/%Y')})"

    # ----- Chụp style dòng mẫu TRƯỚC khi ghi đè bất cứ nội dung nào -----
    template_style = _capture_template_row_style(ws, FIRST_DATA_ROW)

    # ----- Phân nhóm lao động theo đúng thứ tự yêu cầu -----
    groups = {i: [] for i in range(len(GROUP_LABELS))}
    for nv in employees:
        groups[_phan_nhom(nv)].append(nv)

    r = FIRST_DATA_ROW
    stt = 1
    for g_idx, g_label in enumerate(GROUP_LABELS):
        ds_nhom = groups.get(g_idx) or []
        if not ds_nhom:
            continue  # nhóm rỗng -> không in dòng tiêu đề nhóm

        # Sắp xếp trong nhóm theo mã NV tăng dần.
        ds_nhom = sorted(ds_nhom, key=lambda nv: (nv.get("ma_nv") or ""))

        # ----- Dòng tiêu đề nhóm -----
        _write_group_header(ws, r, f"{g_label}: {len(ds_nhom)} người")
        r += 1

        # ----- Các dòng lao động trong nhóm -----
        for nv in ds_nhom:
            _apply_template_row_style(ws, r, template_style)

            dang_dong_bh = bool(nv.get("thang_bat_dau_bh"))

            ws.cell(row=r, column=COL_STT, value=stt)
            ws.cell(row=r, column=COL_MA_NV, value=nv.get("ma_nv") or "")
            ws.cell(row=r, column=COL_HO_TEN, value=(nv.get("ho_ten") or "").strip())
            ws.cell(row=r, column=COL_GIOI_TINH, value=_gioi_tinh_display(nv.get("gioi_tinh")))
            ws.cell(row=r, column=COL_NGAY_SINH, value=_fmt_ngay(nv.get("ngay_sinh")))
            ws.cell(row=r, column=COL_QUOC_TICH, value=nv.get("quoc_tich") or "Việt Nam")
            ws.cell(row=r, column=COL_NOI_CU_TRU, value=nv.get("thuong_tru") or "")
            ws.cell(row=r, column=COL_SO_CCCD, value=nv.get("so_cccd") or "")
            ws.cell(row=r, column=COL_TRINH_DO, value=nv.get("trinh_do") or "")
            ws.cell(row=r, column=COL_BAC_TRINH_DO_NGHE, value=nv.get("bac_trinh_do_nghe") or "")
            ws.cell(row=r, column=COL_VI_TRI_LAM_VIEC, value=nv.get("chuc_danh_nghe") or "")
            ws.cell(row=r, column=COL_LOAI_HDLD, value=nv.get("loai_hop_dong") or "")
            ws.cell(row=r, column=COL_BAT_DAU_LAM_VIEC, value=_fmt_ngay(nv.get("ngay_vao_lam")))
            ws.cell(row=r, column=COL_BHXH, value="x" if dang_dong_bh else "")
            ws.cell(row=r, column=COL_BHYT, value="x" if dang_dong_bh else "")
            ws.cell(row=r, column=COL_BHTN, value="x" if dang_dong_bh else "")

            # Cột Tiền lương - QUAN TRỌNG: phải ghi giá trị SỐ THẬT (int/float),
            # không phải chuỗi, thì Excel mới áp number_format '#,##0' được.
            luong_cell = ws.cell(row=r, column=COL_TIEN_LUONG, value=_to_number(nv.get("luong_bao_hiem")))
            luong_cell.number_format = "#,##0"
            luong_cell.alignment = openpyxl.styles.Alignment(horizontal="right", vertical="center")

            # Cột Chấm dứt HĐLĐ và lý do: chỉ điền khi có ĐỦ CẢ ngày kết thúc
            # VÀ lý do nghỉ; căn giữa ngang/dọc, wrap text, tự giãn chiều cao.
            thoi_diem_cham_dut = ""
            if nv.get("ngay_ket_thuc") and nv.get("ly_do_nghi"):
                thoi_diem_cham_dut = _fmt_ngay(nv.get("ngay_ket_thuc"))
                thoi_diem_cham_dut += f" - {nv.get('ly_do_nghi')}"
            cham_dut_cell = ws.cell(row=r, column=COL_CHAM_DUT_HDLD, value=thoi_diem_cham_dut)
            cham_dut_cell.alignment = openpyxl.styles.Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            _autosize_row_for_wrap(ws, r, COL_CHAM_DUT_HDLD, thoi_diem_cham_dut)

            stt += 1
            r += 1

    wb.save(output_path)
    return output_path


def render_tab_so_qldld(db_engine, format_date, company_config, auto_download_excel):
    """Nội dung tab '📔 Sổ quản lý lao động' — dán vào khối `with t4:`.

    LƯU Ý QUAN TRỌNG: hàm này KHÔNG tự import format_date/COMPANY_CONFIG/
    _auto_download_excel từ app.py. Lý do: khi Streamlit chạy app.py làm
    script chính, app.py được nạp dưới tên module "__main__", KHÔNG phải
    "app" -> nếu viết `from app import ...` bên trong module này, Python sẽ
    không tìm thấy module "app" đã nạp sẵn, và sẽ NẠP LẠI TOÀN BỘ app.py từ
    đầu như một bản chạy song song -> gây lỗi (các lệnh st.* bị gọi trùng).

    Thay vào đó, 3 tham số này phải được TRUYỀN VÀO từ nơi gọi hàm (app.py),
    nơi chúng vốn đã tồn tại sẵn trong cùng file:
        render_tab_so_qldld(
            st.session_state.db_engine,
            format_date=format_date,
            company_config=COMPANY_CONFIG,
            auto_download_excel=_auto_download_excel,
        )
    """
    import streamlit as st
    import pandas as pd
    import psycopg2.extras

    st.subheader("📔 Sổ quản lý lao động")
    st.caption(
        "Theo Nghị định 145/2020/NĐ-CP (Điều 3) - Sổ tích lũy toàn bộ lịch sử lao động "
        "của doanh nghiệp, không phải báo cáo biến động theo kỳ."
    )

    ngay_bao_cao = st.date_input(
        "📅 Tính đến ngày:",
        value=date.today(),
        key="soqldld_ngay_bao_cao",
    )

    col_btn1, col_btn2, col_btn3 = st.columns([1, 2, 1])
    with col_btn2:
        export_clicked = st.button(
            "📥 XUẤT EXCEL SỔ QUẢN LÝ LAO ĐỘNG",
            type="primary",
            width="stretch",
            use_container_width=True,
            key="soqldld_export_btn",
        )

    st.divider()

    # Toàn bộ lao động đã từng vào làm tính đến ngày báo cáo -> đúng nghiệp vụ
    # của Sổ QLLĐ (Điều 3 NĐ 145/2020/NĐ-CP): sổ tích lũy ghi nhận toàn bộ
    # lịch sử lao động, gồm cả người đang làm lẫn người đã chấm dứt HĐLĐ ở BẤT
    # KỲ thời điểm nào trong quá khứ, chứ KHÔNG chỉ những người chấm dứt trong
    # một khoảng thời gian nhất định (đó là nghiệp vụ của Điều 4 - Báo cáo sử
    # dụng lao động / Mẫu 01/PLI, không phải của sổ này).
    db = db_engine.get_connection()
    c = db.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    c.execute(
        """
        SELECT
            nv.ma_nv, nv.ho_ten, nv.gioi_tinh, nv.ngay_sinh, nv.quoc_tich,
            nv.thuong_tru, nv.so_cccd, nv.chuc_danh_nghe, nv.loai_hop_dong,
            nv.ngay_vao_lam, nv.thang_bat_dau_bh, nv.luong_bao_hiem,
            nv.ngay_ket_thuc, nv.ly_do_nghi, nv.trang_thai,
            nv.trinh_do, nv.so_hdld
        FROM nhan_vien nv
        WHERE nv.ngay_vao_lam <= %s
        AND nv.so_hdld IS NOT NULL AND nv.so_hdld != ''
        ORDER BY nv.ngay_vao_lam ASC
        """,
        (ngay_bao_cao,),
    )
    ds_lao_dong = c.fetchall()
    db.close()

    st.markdown(f"### 👥 Danh sách lao động tính đến ngày báo cáo ({len(ds_lao_dong)})")
    if ds_lao_dong:
        df = pd.DataFrame(ds_lao_dong)
        for col in df.columns:
            if "ngay" in col.lower():
                df[col] = df[col].apply(format_date)
        preview_cols = [
            "ma_nv", "ho_ten", "gioi_tinh", "chuc_danh_nghe",
            "loai_hop_dong", "ngay_vao_lam", "ngay_ket_thuc",
        ]
        available_cols = [c for c in preview_cols if c in df.columns]
        df_preview = df[available_cols]
        df_preview.columns = [
            "Mã NV", "Họ tên", "Giới tính", "Vị trí làm việc",
            "Loại HĐLĐ", "Ngày vào làm", "Ngày kết thúc",
        ][: len(available_cols)]
        st.dataframe(df_preview, width="stretch", hide_index=True)
    else:
        st.info("📭 Không có lao động nào được ghi nhận tính đến ngày báo cáo.")

    if export_clicked:
        if ds_lao_dong:
            with st.spinner("Đang tạo Sổ quản lý lao động... Vui lòng chờ..."):
                try:
                    filename = f"SoQLLD_TinhDenNgay_{ngay_bao_cao.strftime('%d%m%Y')}.xlsx"
                    build_so_qldld_excel(
                        template_path=TEMPLATE_SO_QLLD,
                        output_path=filename,
                        employees=ds_lao_dong,
                        ngay_bao_cao=ngay_bao_cao,
                        company_config=company_config,
                    )

                    with open(filename, "rb") as f:
                        file_data = f.read()

                    st.success(f"✅ Đã tạo Sổ quản lý lao động thành công! {len(ds_lao_dong)} lao động.")
                    st.cache_data.clear()

                    auto_download_excel(file_data, filename)

                    if os.path.exists(filename):
                        os.remove(filename)

                except Exception as e:
                    st.error(f"❌ Lỗi khi tạo Sổ quản lý lao động: {str(e)}")
                    st.exception(e)
        else:
            st.warning("⚠️ Không có lao động nào trong kỳ để xuất báo cáo!")
